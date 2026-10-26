"""The resume builder must reproduce the user's gold format, and the validator must pass a clean
resume while rejecting a broken one , so a tailored copy ships without a manual recheck."""
import io

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app import resume_export as RE
from app import resume_validation as RV


SAMPLE = {
    "name": "Krish Ojha", "email": "krish@example.com", "phone": "+91-99999", "location": "Bengaluru",
    "links": ["linkedin.com/in/krishojha", "github.com/KrishOjha1810"],
    "summary": "Backend engineer with nearly 2 years building production systems in Rust and Python.",
    "skills": {"Languages": ["Rust", "Python", "TypeScript"], "Backend & Data": ["FastAPI", "Postgres"]},
    "experience": [{"title": "Software Engineer", "company": "Ancilar", "dates": "Feb 2025 - Present",
                    "bullets": ["Built a 3-tier streaming pipeline.", "Cut runtime errors by 40%."]}],
    "projects": [{"name": "VoiceBridge", "stack": "Claude Code, Whisper",
                  "bullets": ["Voice layer on Claude Code."]}],
    "education": [{"school": "Some Institute", "degree": "B.Tech CSE", "dates": "Nov 2024"}],
}


def _doc(rj):
    return docx.Document(io.BytesIO(RE.build_docx(rj)))


def test_build_docx_matches_gold_geometry_and_fonts():
    d = _doc(SAMPLE)
    s = d.sections[0]
    assert abs(s.page_width.inches - 8.5) < 0.01 and abs(s.page_height.inches - 11) < 0.01
    assert abs(s.top_margin.inches - 0.3) < 0.02 and abs(s.left_margin.inches - 0.45) < 0.02
    assert d.styles["Normal"].font.name == "Cambria"
    assert abs(d.styles["Normal"].font.size.pt - 10) < 0.1


def test_build_docx_name_and_headings():
    d = _doc(SAMPLE)
    paras = [p for p in d.paragraphs if p.text.strip()]
    nm = paras[0]
    assert nm.text == "Krish Ojha"
    assert nm.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert abs(nm.runs[0].font.size.pt - 19) < 0.1 and nm.runs[0].font.bold
    heads = [p.text for p in d.paragraphs
             if p.runs and p.runs[0].bold and p.text.strip().isupper()
             and str(p.runs[0].font.color.rgb) == "1F4E79"]
    assert heads == ["PROFESSIONAL SUMMARY", "SKILLS", "PROFESSIONAL EXPERIENCE", "KEY PROJECTS", "EDUCATION"]


def test_build_docx_bullets_are_literal_glyph_not_word_list():
    d = _doc(SAMPLE)
    bul = [p for p in d.paragraphs if p.runs and p.runs[0].text.startswith("•")]
    assert len(bul) == 3  # 2 experience + 1 project
    assert all(p.style.name == "Normal" for p in bul)  # not "List Bullet"
    assert abs(bul[0].paragraph_format.left_indent.inches - 0.2) < 0.01
    assert bul[0].paragraph_format.first_line_indent.inches < 0  # hanging


def test_validator_passes_clean_resume():
    data = RE.build_docx(SAMPLE)
    res = RV.validate_resume(data, SAMPLE, years=2)  # 1-page tier, no renderer -> estimator
    assert res["ok"] is True, res["failures"]


def test_validator_rejects_broken_resume():
    # a plain python-docx document with none of the gold format / sections
    d = docx.Document()
    d.add_paragraph("Krish Ojha")
    d.add_paragraph("just some text with no sections and no email")
    buf = io.BytesIO(); d.save(buf)
    res = RV.validate_resume(buf.getvalue(), SAMPLE, years=2)
    assert res["ok"] is False
    names_failed = {f.split(":")[0] for f in res["failures"]}
    assert "format" in names_failed and "sections" in names_failed
